import json
from datetime import UTC, datetime
from pathlib import Path

import pytest
from alembic import command
from openpyxl import load_workbook
from sqlalchemy import text

from x_follow_list.application.binding import BindingClaim, BrowserBindingService
from x_follow_list.application.scan_coordination import ScanCoordinator
from x_follow_list.artifacts.xlsx import XlsxArtifactService
from x_follow_list.browser.contracts import BrowserSession, BrowserSessionRequest, ProviderConfig
from x_follow_list.browser.direct_chrome import DirectChromeProvider
from x_follow_list.browser.registry import BrowserProviderRegistry
from x_follow_list.config import RuntimeEnvironment, Settings
from x_follow_list.persistence.database import Database
from x_follow_list.persistence.migrations import alembic_config
from x_follow_list.worker.browser_binding import BrowserBindingJob, DetectedIdentity
from x_follow_list.worker.browser_session import BrowserSessionJob
from x_follow_list.worker.relationship_scan import RelationshipScanJob
from x_follow_list.worker.scan_loop import ScanWorker


async def prepared_database(tmp_path: Path) -> Database:
    settings = Settings(environment=RuntimeEnvironment.TEST, data_dir=tmp_path)
    await command_upgrade(settings)
    database = Database.from_settings(settings)
    now = datetime.now(UTC).isoformat()
    config = json.dumps({"profiles_root": str((tmp_path / "profiles").resolve())})
    async with database.session() as session:
        await session.execute(
            text("INSERT INTO users VALUES ('owner','owner@test','h','OWNER',:n,:n)"),
            {"n": now},
        )
        await session.execute(
            text(
                "INSERT INTO browser_provider_configs "
                "(id,owner_user_id,provider_code,config_version,display_name,config_json,"
                "created_at,updated_at) VALUES "
                "('provider','owner','DIRECT_CHROME',1,'Chrome',:config,:n,:n)"
            ),
            {"config": config, "n": now},
        )
        await session.commit()
    return database


async def command_upgrade(settings: Settings) -> None:
    import asyncio

    await asyncio.to_thread(command.upgrade, alembic_config(settings), "head")


@pytest.mark.asyncio
async def test_direct_chrome_completes_scan_diff_event_and_xlsx_journey(
    tmp_path: Path, fixture_origin: str
) -> None:
    database = await prepared_database(tmp_path)
    coordinator = ScanCoordinator(database)
    artifacts = XlsxArtifactService(database, tmp_path / "artifacts")
    provider = DirectChromeProvider(forbidden_profile_roots=frozenset())
    registry = BrowserProviderRegistry([provider])
    provider_config = ProviderConfig(
        "DIRECT_CHROME", 1, {"profiles_root": str((tmp_path / "profiles").resolve())}
    )
    binding_service = BrowserBindingService(database)
    binding = await binding_service.create("owner", "provider", "phase-a")
    binding_claim = await binding_service.claim_next("release-worker")
    assert binding_claim is not None

    async def binding_request(_claim: BindingClaim) -> BrowserSessionRequest:
        return BrowserSessionRequest(
            provider_config, binding.profile_ref, binding.session_id, headless=False
        )

    async def read_local_identity(session: BrowserSession) -> DetectedIdentity:
        page = session.context.pages[0]
        await page.goto(f"{fixture_origin}/profile?scenario=phase_a_baseline")
        return DetectedIdentity("self", "owner", "Owner")

    await BrowserBindingJob(
        binding_service, registry, binding_request, read_local_identity
    )(binding_claim)
    account_id = await binding_service.confirm("owner", binding.session_id)

    run_ids: list[str] = []
    for index, scenario in enumerate(("phase_a_baseline", "phase_a_changed"), start=1):
        run_id = await coordinator.enqueue(
            "owner", account_id, f"journey-{index}", f"journey-{index}"
        )
        relationship_scan = RelationshipScanJob(
            database,
            artifacts,
            f"{fixture_origin}/profile?scenario={scenario}",
        )
        browser_job = BrowserSessionJob(
            registry,
            lambda current_claim: BrowserSessionRequest(
                provider_config, "phase-a", current_claim.run_id, headless=True
            ),
            relationship_scan,
        )
        assert await ScanWorker(coordinator, "release-worker", browser_job).run_once()
        run_ids.append(run_id)

    async with database.session() as session:
        runs = (
            await session.execute(
                text(
                    "SELECT id,status,follower_count,following_count FROM scan_runs "
                    "ORDER BY created_at,id"
                )
            )
        ).all()
        event = (
            await session.execute(
                text(
                    "SELECT subject_x_user_id,event_type FROM relationship_events "
                    "WHERE scan_run_id=:run AND event_type='UNFOLLOWED_ME_AFTER_MUTUAL'"
                ),
                {"run": run_ids[1]},
            )
        ).one()
        artifact_path = await session.scalar(
            text(
                "SELECT storage_path FROM artifacts "
                "WHERE scan_run_id=:run AND status='READY'"
            ),
            {"run": run_ids[1]},
        )
        active_leases = await session.scalar(
            text("SELECT count(*) FROM resource_leases WHERE expires_at > :now"),
            {"now": datetime.now(UTC).isoformat()},
        )

    workbook = load_workbook(Path(str(artifact_path)), read_only=True)
    unfollowed = list(workbook["UnfollowedMe"].iter_rows(min_row=2, values_only=True))
    workbook.close()
    await database.dispose()

    assert [(row[1], row[2], row[3]) for row in runs] == [
        ("SUCCESS", 2, 2),
        ("SUCCESS", 1, 2),
    ]
    assert event == ("102", "UNFOLLOWED_ME_AFTER_MUTUAL")
    assert unfollowed[0][0] == "102"
    assert active_leases == 0
