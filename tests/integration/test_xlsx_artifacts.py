import asyncio
import hashlib
import json
from datetime import UTC, datetime, timedelta
from pathlib import Path

import httpx
import pytest
from alembic import command
from fastapi import FastAPI
from openpyxl import load_workbook
from pydantic import SecretStr
from sqlalchemy import text

from x_follow_list.api.app import create_app
from x_follow_list.artifacts.xlsx import XlsxArtifactService
from x_follow_list.config import RuntimeEnvironment, Settings
from x_follow_list.persistence.migrations import alembic_config

ORIGIN = "http://testserver"


async def prepared_app(tmp_path: Path) -> FastAPI:
    settings = Settings(
        environment=RuntimeEnvironment.TEST,
        data_dir=tmp_path,
        bootstrap_token=SecretStr("one-time-bootstrap-token"),
        app_origin=ORIGIN,
    )
    await asyncio.to_thread(command.upgrade, alembic_config(settings), "head")
    return create_app(settings)


async def login_and_seed(client: httpx.AsyncClient, app: FastAPI) -> tuple[str, str]:
    bootstrap = await client.post(
        "/api/v1/auth/bootstrap",
        headers={"Origin": ORIGIN},
        json={
            "login": "owner@example.test",
            "password": "a sufficiently long password",
            "bootstrap_token": "one-time-bootstrap-token",
        },
    )
    login = await client.post(
        "/api/v1/auth/session",
        headers={"Origin": ORIGIN},
        json={"login": "owner@example.test", "password": "a sufficiently long password"},
    )
    owner_id = str(bootstrap.json()["id"])
    now = datetime(2026, 9, 15, 4, 0, tzinfo=UTC)
    async with app.state.database.session() as session:
        await session.execute(
            text("UPDATE users SET timezone='Asia/Taipei' WHERE id=:owner"),
            {"owner": owner_id},
        )
        await session.execute(
            text(
                "INSERT INTO browser_provider_configs "
                "(id,owner_user_id,provider_code,config_version,display_name,config_json,"
                "created_at,updated_at) VALUES "
                "('provider',:owner,'DIRECT_CHROME',1,'Chrome',:config,:now,:now)"
            ),
            {"owner": owner_id, "config": json.dumps({}), "now": now.isoformat()},
        )
        await session.execute(
            text(
                "INSERT INTO x_accounts "
                "(id,owner_user_id,provider_config_id,profile_ref,profile_ref_hash,x_user_id,"
                "username,display_name,status,last_successful_scan_at,created_at,updated_at) "
                "VALUES ('account',:owner,'provider','profile','hash','self','owner','Owner',"
                "'READY',:now,:now,:now)"
            ),
            {"owner": owner_id, "now": now.isoformat()},
        )
        await session.execute(
            text(
                "INSERT INTO x_account_memberships "
                "(x_account_id,user_id,role,created_at,updated_at) "
                "VALUES ('account',:owner,'OWNER',:now,:now)"
            ),
            {"owner": owner_id, "now": now.isoformat()},
        )
        await session.execute(
            text(
                "INSERT INTO scan_runs "
                "(id,x_account_id,requested_by_user_id,status,follower_count,following_count,"
                "created_at,finished_at) VALUES "
                "('run', 'account', :owner, 'SUCCESS', 2, 2, :now, :now)"
            ),
            {"owner": owner_id, "now": now.isoformat()},
        )
        await session.execute(
            text(
                "INSERT INTO relationship_snapshots "
                "(id,x_account_id,scan_run_id,follower_count,following_count,captured_at) "
                "VALUES ('snapshot','account','run',2,2,:now)"
            ),
            {"now": now.isoformat()},
        )
        await session.execute(
            text(
                "INSERT INTO snapshot_memberships "
                "(snapshot_id,subject_x_user_id,follows_me,i_follow,username,display_name) VALUES "
                "('snapshot','101',1,1,'mutual',NULL),"
                "('snapshot','102',0,1,'=danger','@Formula'),"
                "('snapshot','103',1,0,'new_follower','New Follower')"
            )
        )
        await session.execute(
            text(
                "INSERT INTO relationship_states "
                "(x_account_id,subject_x_user_id,state,non_followback_streak,last_snapshot_id,"
                "updated_at) VALUES "
                "('account','101','MUTUAL',0,'snapshot',:now),"
                "('account','102','NOT_FOLLOWING_BACK',3,'snapshot',:now),"
                "('account','103','FOLLOWS_ME_ONLY',0,'snapshot',:now)"
            ),
            {"now": now.isoformat()},
        )
        await session.execute(
            text(
                "INSERT INTO relationship_events "
                "(id,x_account_id,scan_run_id,subject_x_user_id,category,event_type,dedupe_key,"
                "status,created_at) VALUES "
                "('unfollowed','account','run','102','ACTION_ITEM',"
                "'UNFOLLOWED_ME_AFTER_MUTUAL','d1','OPEN',:now),"
                "('new','account','run','103','RELATIONSHIP_CHANGE',"
                "'NEW_FOLLOWER','d2','INFORMATIONAL',:now)"
            ),
            {"now": now.isoformat()},
        )
        await session.commit()
    return owner_id, str(login.json()["csrf_token"])


def headers(csrf: str) -> dict[str, str]:
    return {"Origin": ORIGIN, "X-CSRF-Token": csrf}


@pytest.mark.asyncio
async def test_builds_six_snapshot_pinned_sheets_and_downloads_with_safe_headers(
    tmp_path: Path,
) -> None:
    app = await prepared_app(tmp_path)
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url=ORIGIN
    ) as client:
        _owner, csrf = await login_and_seed(client, app)
        created = await client.post(
            "/api/v1/scan-runs/run/artifacts/xlsx", headers=headers(csrf)
        )
        replay = await client.post(
            "/api/v1/scan-runs/run/artifacts/xlsx", headers=headers(csrf)
        )
        artifact_id = str(created.json()["id"])
        downloaded = await client.get(f"/api/v1/artifacts/{artifact_id}/download")

    assert created.status_code == replay.status_code == 200
    assert created.json()["id"] == replay.json()["id"]
    assert created.json()["status"] == "READY"
    assert created.json()["sha256"] == hashlib.sha256(downloaded.content).hexdigest()
    assert created.json()["byte_size"] == len(downloaded.content)
    assert created.json()["expires_at"]
    assert created.json()["deleted_at"] is None
    assert downloaded.status_code == 200
    assert downloaded.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert downloaded.headers["x-content-type-options"] == "nosniff"
    assert downloaded.headers["content-disposition"].startswith("attachment;")

    path = tmp_path / "downloaded.xlsx"
    path.write_bytes(downloaded.content)
    workbook = load_workbook(path, read_only=True, data_only=False)
    assert workbook.sheetnames == [
        "Summary",
        "UnfollowedMe",
        "NonFollowers",
        "NewFollowers",
        "Followers",
        "Following",
    ]
    summary: dict[object, object] = {
        row[0]: row[1]
        for row in workbook["Summary"].iter_rows(min_row=2, values_only=True)
    }
    assert summary["scan_run_id"] == "run"
    assert summary["snapshot_id"] == "snapshot"
    assert summary["captured_at_utc"] == "2026-09-15T04:00:00Z"
    assert summary["captured_at_local"] == "2026-09-15T12:00:00+08:00"
    non_follower = list(workbook["NonFollowers"].iter_rows(min_row=2, values_only=True))[0]
    assert non_follower[0:3] == ("102", "'=danger", "'@Formula")
    assert non_follower[3:5] == ("NOT_FOLLOWING_BACK", 3)
    follower_rows = list(workbook["Followers"].iter_rows(min_row=2, values_only=True))
    assert {row[0] for row in follower_rows} == {"101", "103"}
    assert next(row for row in follower_rows if row[0] == "101")[2] is None
    workbook.close()
    await app.state.database.dispose()


@pytest.mark.asyncio
async def test_membership_download_authorization_hides_foreign_and_deleted_artifacts(
    tmp_path: Path,
) -> None:
    app = await prepared_app(tmp_path)
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url=ORIGIN
    ) as client:
        owner_id, csrf = await login_and_seed(client, app)
        created = await client.post(
            "/api/v1/scan-runs/run/artifacts/xlsx", headers=headers(csrf)
        )
        artifact_id = str(created.json()["id"])
        now = datetime.now(UTC)
        async with app.state.database.session() as session:
            await session.execute(
                text(
                    "INSERT INTO users "
                    "(id,email,password_hash,role,timezone,created_at,updated_at) VALUES "
                    "('viewer','viewer@example.test','unused','VIEWER','UTC',:now,:now),"
                    "('foreign','foreign@example.test','unused','OWNER','UTC',:now,:now)"
                ),
                {"now": now.isoformat()},
            )
            for user_id in ("viewer", "foreign"):
                await session.execute(
                    text(
                        "INSERT INTO auth_sessions "
                        "(id,user_id,token_hash,csrf_token_hash,created_at,expires_at,"
                        "last_seen_at) "
                        "VALUES (:id,:user,:token,:csrf,:now,:expires,:now)"
                    ),
                    {
                        "id": f"{user_id}-session",
                        "user": user_id,
                        "token": hashlib.sha256(f"{user_id}-token".encode()).hexdigest(),
                        "csrf": hashlib.sha256(b"csrf").hexdigest(),
                        "now": now.isoformat(),
                        "expires": (now + timedelta(hours=1)).isoformat(),
                    },
                )
            await session.execute(
                text(
                    "INSERT INTO x_account_memberships "
                    "(x_account_id,user_id,role,created_at,updated_at) "
                    "VALUES ('account','viewer','VIEWER',:now,:now)"
                ),
                {"now": now.isoformat()},
            )
            await session.commit()

        client.cookies.set("x_follow_list_session", "viewer-token")
        viewer = await client.get(f"/api/v1/artifacts/{artifact_id}/download")
        client.cookies.set("x_follow_list_session", "foreign-token")
        foreign = await client.get(f"/api/v1/artifacts/{artifact_id}/download")
        client.cookies.clear()
        login = await client.post(
            "/api/v1/auth/session",
            headers={"Origin": ORIGIN},
            json={"login": "owner@example.test", "password": "a sufficiently long password"},
        )
        assert owner_id
        async with app.state.database.session() as session:
            await session.execute(
                text("UPDATE artifacts SET status='DELETED',deleted_at=:now WHERE id=:id"),
                {"now": now.isoformat(), "id": artifact_id},
            )
            await session.commit()
        client.cookies.set("x_follow_list_session", login.cookies["x_follow_list_session"])
        deleted = await client.get(f"/api/v1/artifacts/{artifact_id}/download")

    await app.state.database.dispose()
    assert viewer.status_code == 200
    assert foreign.status_code == deleted.status_code == 404


@pytest.mark.asyncio
async def test_failed_publish_leaves_no_partial_file_and_can_be_rebuilt(
    tmp_path: Path,
) -> None:
    app = await prepared_app(tmp_path)
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url=ORIGIN
    ) as client:
        _owner, _csrf = await login_and_seed(client, app)

    service = XlsxArtifactService(app.state.database, tmp_path / "artifacts")
    with pytest.raises(OSError, match="injected"):
        await service.build("run", fail_after_rows=2)

    assert not list((tmp_path / "artifacts").rglob("*.tmp"))
    assert not list((tmp_path / "artifacts").rglob("*.xlsx"))
    async with app.state.database.session() as session:
        failed = (
            await session.execute(text("SELECT status,storage_path FROM artifacts"))
        ).one()
    assert failed[0] == "FAILED"

    rebuilt = await service.build("run")
    assert rebuilt["status"] == "READY"
    assert await asyncio.to_thread(Path(str(rebuilt["storage_path"])).is_file)
    await app.state.database.dispose()
