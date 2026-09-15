from pathlib import Path
from tracemalloc import get_traced_memory, start, stop

from openpyxl import load_workbook

from x_follow_list.artifacts.xlsx import sanitize_excel_text, write_xlsx_stream


def test_external_text_is_never_interpreted_as_an_excel_formula() -> None:
    for value in ("=cmd|' /C calc'!A0", "+SUM(1,1)", "-1+2", "@evil"):
        assert sanitize_excel_text(value) == f"'{value}"
    assert sanitize_excel_text(" safe") == " safe"
    assert sanitize_excel_text(None) is None


def test_write_only_workbook_keeps_fifty_thousand_rows_below_memory_budget(
    tmp_path: Path,
) -> None:
    target = tmp_path / "large.xlsx"

    def rows():
        yield ("x_user_id", "username", "display_name")
        for index in range(50_000):
            yield (str(index), f"user-{index}", None)

    start()
    write_xlsx_stream(target, (("Followers", rows()),))
    _, peak = get_traced_memory()
    stop()

    workbook = load_workbook(target, read_only=True)
    assert workbook["Followers"].max_row == 50_001
    workbook.close()
    assert peak < 128 * 1024 * 1024
